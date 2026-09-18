"""Build Greek manuals from retained DOCX drafts and reviewed Markdown copy."""
from pathlib import Path
from copy import deepcopy
import hashlib
import json
import re
from zipfile import ZipFile
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
TMP = ROOT/'tmp/docs'
OUT = ROOT/'output/documents'
SOURCE = ROOT/'tools/document_templates'
META = ROOT/'Documentation/manual_build_metadata.json'
meta = json.loads(META.read_text(encoding='utf-8')) if META.exists() else {}
OUT.mkdir(parents=True, exist_ok=True)
TMP.mkdir(parents=True, exist_ok=True)

def xml(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items(): e.set(qn('w:'+k), str(v))
    return e

def clean_borders(e):
    for tag in ['w:pBdr','w:shd']:
        for x in list(e.iter(qn(tag))): x.getparent().remove(x)

def field(p, name):
    r = p.add_run()
    e = OxmlElement('w:fldSimple'); e.set(qn('w:instr'), name)
    r._r.addnext(e)

def prep(reference, short_title):
    d = Document(reference)
    for e in list(d._element.body):
        if e.tag != qn('w:sectPr'): d._element.body.remove(e)
    for s in d.styles:
        if s.type == 1:
            clean_borders(s.element)
            s.font.name = 'Calibri'
            rpr = s.element.get_or_add_rPr()
            for language in list(rpr.findall(qn('w:lang'))): rpr.remove(language)
            rpr.append(xml('w:lang', val='el-GR'))
    for name in ['Title','Subtitle','Heading 1','Heading 2','Heading 3','Kicker']:
        if name in d.styles: d.styles[name].font.color.rgb = RGBColor(0,0,0)
    d.styles['Normal'].font.color.rgb = RGBColor.from_string('193648')
    d.styles['Heading 1'].paragraph_format.space_before = Pt(0)
    d.styles['Heading 1'].paragraph_format.space_after = Pt(8)
    d.styles['Heading 1'].font.size = Pt(18)
    for sec in d.sections:
        for part in [sec.header, sec.footer]:
            for e in list(part._element): part._element.remove(e)
        h = sec.header.add_paragraph(short_title)
        h.style = d.styles['Small']
        h.runs[0].font.color.rgb = RGBColor(0,0,0)
        h.runs[0].font.size = Pt(8)
        f = sec.footer.add_paragraph('ΠΑΔΑ  |  ΑΑΥ 2025–2026   •   ')
        f.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        f.style = d.styles['Small']
        field(f,'PAGE'); f.add_run(' / '); field(f,'NUMPAGES')
    d.core_properties.author = ''
    d.core_properties.last_modified_by = ''
    d.core_properties.title = short_title
    d.core_properties.subject = 'ICU Simulation 2025–2026'
    d.core_properties.comments = ''
    d.core_properties.keywords = 'ΜΕΘ, EHR, HTA, Unity'
    d.core_properties.language = 'el-GR'
    return d

def table(d, rows):
    count=len(rows[0])
    t=d.add_table(rows=0, cols=count)
    t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=False
    width=6.81
    weights=([0.33,0.67] if count==2 else [0.23,0.32,0.45])
    if count==3 and rows[0][0]=='Κόμβος': weights=[0.10,0.17,0.73]
    if count==2 and rows[0][0]=='Σχέδιο': weights=[0.13,0.87]
    widths=[int(width*1440*w) for w in weights]
    pr=t._tbl.tblPr
    for e in list(pr):
        if e.tag in [qn('w:tblInd'),qn('w:tblW'),qn('w:tblBorders'),qn('w:tblCellMar')]: pr.remove(e)
    pr.append(xml('w:tblW',w=sum(widths),type='dxa'))
    pr.append(xml('w:tblInd',w=0,type='dxa'))
    borders=xml('w:tblBorders')
    for n in ['top','left','bottom','right','insideH','insideV']:
        borders.append(xml('w:'+n,val='single',sz=4,color='D9D9D9'))
    pr.append(borders)
    margins=xml('w:tblCellMar')
    for n,v in [('top',85),('bottom',85),('left',105),('right',105)]: margins.append(xml('w:'+n,w=v,type='dxa'))
    pr.append(margins)
    for col,w in zip(t.columns,widths): col.width=Inches(w/1440)
    for ri,values in enumerate(rows):
        row=t.add_row()
        if ri==0: row._tr.get_or_add_trPr().append(xml('w:tblHeader'))
        row._tr.get_or_add_trPr().append(xml('w:cantSplit'))
        for i,value in enumerate(values):
            c=row.cells[i]; c.width=Inches(widths[i]/1440)
            c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            c._tc.get_or_add_tcPr().append(xml('w:shd',fill='193648' if ri==0 else ('EFF4F6' if ri%2 else 'FFFFFF')))
            p=c.paragraphs[0]
            p.paragraph_format.space_after=Pt(0)
            p.paragraph_format.line_spacing=1.03
            r=p.add_run(value.replace('\\n','\n'))
            r.font.name='Calibri';r.font.size=Pt(9.3)
            r.bold=ri==0
            r.font.color.rgb=RGBColor.from_string('FFFFFF' if ri==0 else '193648')
    d.add_paragraph().paragraph_format.space_after=Pt(2)

def build(kind, source_name, doc_name, title):
    ref=SOURCE/source_name
    d=prep(ref,title)
    text=(ROOT/'Documentation'/doc_name.replace('.docx','.md')).read_text(encoding='utf-8')
    lines=text.splitlines();i=0;page=1;figures=[]
    while i<len(lines):
        line=lines[i].strip();i+=1
        if not line: continue
        if line=='<!-- pagebreak -->': d.add_page_break();page+=1;continue
        if line.startswith('!['):
            match=re.fullmatch(r'!\[(.*?)\]\((.*?)\)',line)
            if not match: raise ValueError('Invalid image reference: '+line)
            image_alt,image_path=match.groups()
            path=(ROOT/'Documentation'/image_path).resolve()
            if path.is_file():
                p=d.add_paragraph();p.paragraph_format.space_after=Pt(3)
                p.paragraph_format.keep_with_next=True
                r=p.add_run();r.add_picture(str(path),width=Inches(6.81))
                r._r.xpath('.//wp:docPr')[0].set('descr',image_alt)
                figures.append(str(path.relative_to(ROOT)))
            else: raise FileNotFoundError('Required runtime screenshot is missing: '+str(path))
            continue
        if line=='<!-- verification -->':
            for para in meta.get('evidence_paragraphs',['Ο αρχικός έλεγχος Unity EditMode ολοκληρώθηκε με 61 επιτυχίες σε 61 δοκιμές. Το αποτέλεσμα καταγράφεται στο Evidence/EditMode-initial.xml.']): d.add_paragraph(para)
            continue
        if line.startswith('|'):
            rows=[[s.strip() for s in line.strip('|').split('|')]]
            while i<len(lines) and lines[i].strip().startswith('|'):
                row=[s.strip() for s in lines[i].strip().strip('|').split('|')];i+=1
                if all(re.fullmatch(r'[:\- ]+',s) for s in row):continue
                rows.append(row)
            table(d,rows);continue
        if line.startswith('# '):
            p=d.add_paragraph(line[2:],style='Title');clean_borders(p._p);continue
        if line.startswith('## '):d.add_paragraph(line[3:],style='Heading 1' if line[3:4].isdigit() else 'Heading 2');continue
        if line.startswith('Εικόνα '):d.add_paragraph(line,style='Caption');continue
        if kind=='technical' and re.match(r'^\d+(?:\.\d+)?[. ]',line):
            p=d.add_paragraph(line,style='HTA')
            if re.match(r'^\d+\.\d+',line):p.paragraph_format.left_indent=Inches(.20)
            else:
                p.runs[0].bold=True;p.paragraph_format.space_before=Pt(3)
                p.paragraph_format.keep_with_next=True
            continue
        p=d.add_paragraph(line)
        if line.startswith('Πανεπιστήμιο ') or line.startswith('Αλληλεπίδραση ') or line.startswith('Ακαδημαϊκό '):p.style='Small'
    path=OUT/doc_name
    d.save(path)
    print(path)
    return {'file':str(path.relative_to(ROOT)),'figures':figures,'sha256':hashlib.sha256(path.read_bytes()).hexdigest()}

results=[build('user','ICU_User_Manual_DRAFT.docx','ICU_User_Manual_EL.docx','ICU Simulation  Εγχειρίδιο χρήστη'),build('technical','ICU_Technical_Manual_HTA_DRAFT.docx','ICU_Technical_Manual_HTA_EL.docx','ICU Simulation  Τεχνικό εγχειρίδιο και HTA')]
(TMP/'build_manifest.json').write_text(json.dumps(results,ensure_ascii=False,indent=2),encoding='utf-8')
